"""
StampCatalogue Application
--------------------------
A standalone desktop application for managing multiple stamp collections,
searching country catalogues, and exporting Excel mancolists.

Reuses the exact same image resolution logic (getPhotoName) and localization
framework (gettext + AppLanguage) as StampAlbum.

Author: Boris du Reau / AI Assistant
Year: 2026
"""

import sys
import os
import sqlite3
import gettext
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QComboBox, QLabel, QTableWidget, QTableWidgetItem, QPushButton,
    QHeaderView, QLineEdit, QSpinBox, QDoubleSpinBox, QMessageBox,
    QFileDialog, QGroupBox, QFormLayout, QTextEdit, QSplitter,
    QDialog, QDialogButtonBox
)
from PyQt6.QtGui import QPixmap

# Localization setup following StampAlbum's architecture
from AppLanguage import get_language

gettext.find("StampCatalogue")
translate = gettext.translation(
    'StampCatalogue', localedir='locale', languages=[get_language()], fallback=True
)
translate.install()
_ = translate.gettext


class CollectionManagerDB:
    """Manages the central 'collections.db' storing user collections and owned stamps."""

    def __init__(self, db_dir: Path):
        self.db_path = db_dir / "collections.db"
        self.conn = sqlite3.connect(self.db_path)
        self.cursor = self.conn.cursor()
        self._init_db()

    def _init_db(self):
        """Initialize central tables if they do not exist."""
        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS collections (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                country TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
        """)
        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS collection_items (
                collection_id INTEGER NOT NULL,
                stamp_nbr TEXT NOT NULL,
                status TEXT NOT NULL,
                condition TEXT,
                quantity INTEGER DEFAULT 1,
                purchase_price REAL DEFAULT 0.0,
                notes TEXT,
                PRIMARY KEY (collection_id, stamp_nbr),
                FOREIGN KEY (collection_id) REFERENCES collections(id) ON DELETE CASCADE
            );
        """)
        self.conn.commit()

    def get_collections(self):
        """Fetch all user collections."""
        self.cursor.execute("SELECT id, name, country FROM collections ORDER BY name ASC")
        return self.cursor.fetchall()

    def create_collection(self, name: str, country: str) -> int:
        """Create a new collection linked to a specific country."""
        self.cursor.execute("INSERT INTO collections (name, country) VALUES (?, ?)", (name, country))
        self.conn.commit()
        return self.cursor.lastrowid

    def delete_collection(self, collection_id: int):
        """Delete a collection and all associated stamp items."""
        self.cursor.execute("DELETE FROM collections WHERE id = ?", (collection_id,))
        self.cursor.execute("DELETE FROM collection_items WHERE collection_id = ?", (collection_id,))
        self.conn.commit()

    def save_item(self, collection_id: int, stamp_nbr: str, status: str, condition: str, quantity: int, price: float,
                  notes: str):
        """Insert or update a stamp record inside a collection."""
        query = """
            INSERT INTO collection_items (collection_id, stamp_nbr, status, condition, quantity, purchase_price, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(collection_id, stamp_nbr) DO UPDATE SET
                status = excluded.status,
                condition = excluded.condition,
                quantity = excluded.quantity,
                purchase_price = excluded.purchase_price,
                notes = excluded.notes;
        """
        self.cursor.execute(query, (collection_id, str(stamp_nbr), status, condition, quantity, price, notes))
        self.conn.commit()

    def remove_item(self, collection_id: int, stamp_nbr: str):
        """Remove a stamp from a collection."""
        self.cursor.execute("DELETE FROM collection_items WHERE collection_id = ? AND stamp_nbr = ?",
                            (collection_id, str(stamp_nbr)))
        self.conn.commit()

    def get_collection_items(self, collection_id: int):
        """Fetch all saved items for a given collection ID."""
        self.cursor.execute("""
            SELECT stamp_nbr, status, condition, quantity, purchase_price, notes 
            FROM collection_items WHERE collection_id = ?
        """, (collection_id,))
        return {row[0]: row[1:] for row in self.cursor.fetchall()}

    def close(self):
        """Close database connection."""
        if self.conn:
            self.conn.close()


class CountryCatalogueDB:
    """Read-only database manager for country-specific stamp catalogues."""

    def __init__(self, country: str, db_dir: Path):
        self.country = country
        self.db_path = db_dir / f"{country}.db"

        # Open in read-only mode using SQLite URI
        uri = f"file:{self.db_path.as_posix()}?mode=ro"
        self.conn = sqlite3.connect(uri, uri=True)
        self.cursor = self.conn.cursor()

    def load_types(self):
        """Load unique categories/types for the current country."""
        self.cursor.execute("SELECT DISTINCT type FROM Stamp_List WHERE type IS NOT NULL ORDER BY type ASC")
        return [row[0] for row in self.cursor.fetchall()]

    def load_years(self):
        """Load unique issue years for the current country."""
        self.cursor.execute("SELECT DISTINCT year FROM Stamp_List WHERE year IS NOT NULL ORDER BY year ASC")
        return [row[0] for row in self.cursor.fetchall()]

    def get_stamps(self, stamp_type=None, year=None, search_nbr=None):
        """Query stamps from catalogue with optional filters and sub_nbr field."""
        query = """
            SELECT nbr, year, type, valuecolor, stampDescription, key, sub_nbr
            FROM Stamp_List
            WHERE 1=1
        """
        params = []

        if stamp_type and stamp_type != _("All"):
            query += " AND type = ?"
            params.append(stamp_type)

        if year and year != _("All"):
            query += " AND year = ?"
            params.append(year)

        if search_nbr:
            query += " AND nbr LIKE ?"
            params.append(f"%{search_nbr}%")

        query += " ORDER BY sequence, ascii_seq, nbr ASC"
        self.cursor.execute(query, params)
        return self.cursor.fetchall()

    def close(self):
        """Close database connection."""
        if self.conn:
            self.conn.close()


class CreateCollectionDialog(QDialog):
    """Dialog for creating a new collection associated with a country."""

    def __init__(self, available_countries, parent=None):
        super().__init__(parent)
        self.setWindowTitle(_("Create New Collection"))
        self.resize(350, 150)

        layout = QVBoxLayout(self)
        form = QFormLayout()

        self.edit_name = QLineEdit()
        self.edit_name.setPlaceholderText(_("e.g. France Mint 1900-1950"))
        form.addRow(_("Collection Name:"), self.edit_name)

        self.combo_country = QComboBox()
        self.combo_country.addItems(available_countries)
        form.addRow(_("Country:"), self.combo_country)

        layout.addLayout(form)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def get_data(self):
        return self.edit_name.text().strip(), self.combo_country.currentText()


class MainWindow(QMainWindow):
    """Main Application Window for StampCatalogue."""

    def __init__(self):
        super().__init__()
        self.setWindowTitle(_("StampCatalogue - Stamp Collection Manager"))
        self.resize(1200, 750)

        self.base_dir = self._get_app_dir()
        self.db_dir = self.base_dir / "databases"

        self.central_db = CollectionManagerDB(self.db_dir)
        self.country_db = None

        self.current_collection_id = None
        self.current_country = None
        self.selected_nbr = None

        self.available_countries = self._detect_available_countries()

        self._init_ui()
        self._load_collections()

    def _get_app_dir(self) -> Path:
        if getattr(sys, 'frozen', False):
            return Path(sys.executable).resolve().parent
        return Path(__file__).resolve().parent

    def _detect_available_countries(self):
        """Detect country database files in the databases/ folder."""
        if not self.db_dir.exists():
            return []
        return sorted([f.stem for f in self.db_dir.glob("*.db") if f.stem not in ["master", "collections"]])

    def _init_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        main_layout = QVBoxLayout(main_widget)

        # Collection Header Bar
        coll_box = QGroupBox(_("Collection Selection"))
        coll_layout = QHBoxLayout()

        coll_layout.addWidget(QLabel(_("Active Collection:")))
        self.combo_collections = QComboBox()
        self.combo_collections.currentIndexChanged.connect(self._on_collection_changed)
        coll_layout.addWidget(self.combo_collections, stretch=1)

        btn_new_coll = QPushButton(_("New Collection..."))
        btn_new_coll.clicked.connect(self._create_collection)
        coll_layout.addWidget(btn_new_coll)

        btn_del_coll = QPushButton(_("Delete Collection"))
        btn_del_coll.clicked.connect(self._delete_collection)
        coll_layout.addWidget(btn_del_coll)

        btn_export = QPushButton(_("Export Mancolist (Excel)"))
        btn_export.setStyleSheet("background-color: #1d6f42; color: white; font-weight: bold;")
        btn_export.clicked.connect(self._export_mancolist)
        coll_layout.addWidget(btn_export)

        coll_box.setLayout(coll_layout)
        main_layout.addWidget(coll_box)

        # Filter Bar
        filter_box = QGroupBox(_("Catalogue Filters"))
        filter_layout = QHBoxLayout()

        filter_layout.addWidget(QLabel(_("Type:")))
        self.combo_type = QComboBox()
        self.combo_type.currentTextChanged.connect(self._refresh_table)
        filter_layout.addWidget(self.combo_type)

        filter_layout.addWidget(QLabel(_("Year:")))
        self.combo_year = QComboBox()
        self.combo_year.currentTextChanged.connect(self._refresh_table)
        filter_layout.addWidget(self.combo_year)

        filter_layout.addWidget(QLabel(_("Catalog N°:")))
        self.edit_search_nbr = QLineEdit()
        self.edit_search_nbr.setPlaceholderText(_("e.g. 154..."))
        self.edit_search_nbr.textChanged.connect(self._refresh_table)
        filter_layout.addWidget(self.edit_search_nbr)

        filter_box.setLayout(filter_layout)
        main_layout.addWidget(filter_box)

        # Main Splitter
        splitter = QSplitter(Qt.Orientation.Horizontal)

        # Stamps Table
        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels([
            _("Yvert N°"),
            _("Year"),
            _("Type"),
            _("Value / Color"),
            _("Description"),
            _("Status")
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.itemSelectionChanged.connect(self._on_stamp_selected)
        splitter.addWidget(self.table)

        # Right Panel (Details + Collection Form)
        right_panel = QWidget()
        right_layout = QVBoxLayout(right_panel)

        # Image Viewer
        self.img_label = QLabel(_("No Photo Available"))
        self.img_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.img_label.setFixedSize(220, 220)
        self.img_label.setStyleSheet("border: 1px solid #cbd5e0; background: #f7fafc;")
        right_layout.addWidget(self.img_label, alignment=Qt.AlignmentFlag.AlignCenter)

        # Details Box
        detail_box = QGroupBox(_("Stamp Details"))
        detail_form = QFormLayout()

        self.lbl_valeur = QLabel("-")
        self.lbl_desc = QLabel("-")
        self.lbl_desc.setWordWrap(True)

        detail_form.addRow(_("Value / Color:"), self.lbl_valeur)
        detail_form.addRow(_("Description:"), self.lbl_desc)
        detail_box.setLayout(detail_form)
        right_layout.addWidget(detail_box)

        # Ownership Form
        form_box = QGroupBox(_("Collection Status"))
        form = QFormLayout()

        self.combo_status = QComboBox()
        self.combo_status.addItems([
            _("Not Owned"),
            _("Owned"),
            _("Manco (Wanted)"),
            _("Ordered")
        ])

        self.combo_condition = QComboBox()
        self.combo_condition.addItems([
            _("Mint Never Hinged (** MNH)"),
            _("Mint Hinged (* MH)"),
            _("Used (Cancelled)"),
            _("On Cover")
        ])

        self.spin_qty = QSpinBox()
        self.spin_qty.setRange(1, 100)

        self.spin_price = QDoubleSpinBox()
        self.spin_price.setRange(0, 100000)
        self.spin_price.setSuffix(" €")

        self.txt_notes = QTextEdit()
        self.txt_notes.setMaximumHeight(60)

        form.addRow(_("Status:"), self.combo_status)
        form.addRow(_("Condition:"), self.combo_condition)
        form.addRow(_("Quantity:"), self.spin_qty)
        form.addRow(_("Purchase Price:"), self.spin_price)
        form.addRow(_("Notes:"), self.txt_notes)

        btn_save = QPushButton(_("Save to Collection"))
        btn_save.setStyleSheet("background-color: #2b6cb0; color: white; font-weight: bold;")
        btn_save.clicked.connect(self._save_stamp_status)
        form.addRow(btn_save)

        form_box.setLayout(form)
        right_layout.addWidget(form_box)

        splitter.addWidget(right_panel)
        splitter.setSizes([650, 350])

        main_layout.addWidget(splitter)

    def _load_collections(self):
        """Populate the collections dropdown."""
        self.combo_collections.blockSignals(True)
        self.combo_collections.clear()

        collections = self.central_db.get_collections()
        for coll_id, name, country in collections:
            display_text = f"{name} ({country})"
            self.combo_collections.addItem(display_text, userData={"id": coll_id, "country": country})

        self.combo_collections.blockSignals(False)

        if collections:
            self._on_collection_changed(0)
        else:
            self._clear_ui()

    def _create_collection(self):
        """Open dialog to create a new collection."""
        if not self.available_countries:
            QMessageBox.warning(self, _("Warning"), _("No country databases found in 'databases/'."))
            return

        dialog = CreateCollectionDialog(self.available_countries, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            name, country = dialog.get_data()
            if not name:
                QMessageBox.warning(self, _("Error"), _("Collection name cannot be empty."))
                return
            try:
                coll_id = self.central_db.create_collection(name, country)
                self._load_collections()
                # Select newly created collection
                idx = self.combo_collections.findData({"id": coll_id, "country": country})
                if idx != -1:
                    self.combo_collections.setCurrentIndex(idx)
            except sqlite3.IntegrityError:
                QMessageBox.critical(self, _("Error"), _("A collection with this name already exists."))

    def _delete_collection(self):
        """Delete current selected collection."""
        if self.current_collection_id is None:
            return

        reply = QMessageBox.question(
            self,
            _("Delete Collection"),
            _("Are you sure you want to delete this collection and all its tracked stamps?"),
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if reply == QMessageBox.StandardButton.Yes:
            self.central_db.delete_collection(self.current_collection_id)
            self._load_collections()

    def _on_collection_changed(self, index):
        data = self.combo_collections.currentData()
        if not data:
            self._clear_ui()
            return

        self.current_collection_id = data["id"]
        self.current_country = data["country"]

        # Re-open country database connection in read-only mode
        if self.country_db:
            self.country_db.close()
        self.country_db = CountryCatalogueDB(self.current_country, self.db_dir)

        # Update filters
        self.combo_type.blockSignals(True)
        self.combo_type.clear()
        self.combo_type.addItem(_("All"))
        self.combo_type.addItems(self.country_db.load_types())
        self.combo_type.blockSignals(False)

        self.combo_year.blockSignals(True)
        self.combo_year.clear()
        self.combo_year.addItem(_("All"))
        self.combo_year.addItems([str(y) for y in self.country_db.load_years()])
        self.combo_year.blockSignals(False)

        self._refresh_table()

    def _refresh_table(self):
        if not self.country_db or self.current_collection_id is None:
            return

        stamps = self.country_db.get_stamps(
            stamp_type=self.combo_type.currentText(),
            year=self.combo_year.currentText(),
            search_nbr=self.edit_search_nbr.text().strip()
        )

        user_items = self.central_db.get_collection_items(self.current_collection_id)

        self.table.setRowCount(0)
        for row_idx, stamp in enumerate(stamps):
            nbr, year, stype, valcolor, desc, key, sub_nbr = stamp
            self.table.insertRow(row_idx)

            self.table.setItem(row_idx, 0, QTableWidgetItem(str(nbr or "")))
            self.table.setItem(row_idx, 1, QTableWidgetItem(str(year or "")))
            self.table.setItem(row_idx, 2, QTableWidgetItem(str(stype or "")))
            self.table.setItem(row_idx, 3, QTableWidgetItem(str(valcolor or "")))
            self.table.setItem(row_idx, 4, QTableWidgetItem(str(desc or "")))

            # User possession status
            user_data = user_items.get(str(nbr))
            status_text = user_data[0] if user_data else _("Not Owned")
            status_item = QTableWidgetItem(status_text)

            if status_text == _("Owned"):
                status_item.setForeground(Qt.GlobalColor.darkGreen)
            elif status_text == _("Manco (Wanted)"):
                status_item.setForeground(Qt.GlobalColor.red)

            self.table.setItem(row_idx, 5, status_item)

            full_data = stamp + (user_data,)
            self.table.item(row_idx, 0).setData(Qt.ItemDataRole.UserRole, full_data)

    def _on_stamp_selected(self):
        selected_rows = self.table.selectedItems()
        if not selected_rows:
            return

        row = self.table.currentRow()
        stamp_data = self.table.item(row, 0).data(Qt.ItemDataRole.UserRole)

        nbr, year, stype, valcolor, desc, key, sub_nbr, user_data = stamp_data
        self.selected_nbr = nbr

        self.lbl_valeur.setText(str(valcolor or "-"))
        self.lbl_desc.setText(str(desc or "-"))

        if user_data:
            status, cond, qty, price, notes = user_data
            self.combo_status.setCurrentText(status)
            self.combo_condition.setCurrentText(cond if cond else _("Mint Never Hinged (** MNH)"))
            self.spin_qty.setValue(qty if qty else 1)
            self.spin_price.setValue(price if price else 0.0)
            self.txt_notes.setPlainText(notes if notes else "")
        else:
            self.combo_status.setCurrentText(_("Not Owned"))
            self.combo_condition.setCurrentIndex(0)
            self.spin_qty.setValue(1)
            self.spin_price.setValue(0.0)
            self.txt_notes.clear()

        # Load image using exact getPhotoName logic from Stamp Album
        self._load_stamp_image(nbr, stype, sub_nbr)

    def getPhotoName(self, stampNumber: str, stampType: str, stampSubNbr) -> str:
        """Generate relative photo image file name based on stamp category prefix (exact match from Stamp Album)."""
        subNbr = ""

        if stampSubNbr is not None and str(stampSubNbr).strip() not in ["", "None"]:
            subNbr = "-" + str(stampSubNbr)

        stampNumber_str = str(stampNumber)
        type_prefix = stampNumber_str[0:2]

        # Aérien
        if type_prefix == "A ":
            nbr = "T01-030-" + stampNumber_str[2:] + subNbr + ".jpg"
        # Bloc feuillet
        elif type_prefix == "BF":
            nbr = "T01-080-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Service
        elif type_prefix == "S ":
            nbr = "T01-170-" + stampNumber_str[2:] + subNbr + ".jpg"
        # bloc carnet
        elif type_prefix == "BC":
            nbr = "T01-000-BC" + stampNumber_str[3:] + subNbr + ".jpg"
        # bande
        elif type_prefix == "B ":
            nbr = "T01-000-B" + stampNumber_str[2:] + subNbr + ".jpg"
        # paire
        elif type_prefix == "P ":
            nbr = "T01-000-P" + stampNumber_str[2:] + subNbr + ".jpg"
        # préoblitéré
        elif type_prefix == "PR":
            nbr = "T01-420-" + stampNumber_str[3:] + subNbr + ".jpg"
        # franchise militaire
        elif type_prefix == "FM":
            nbr = "T01-190-" + stampNumber_str[3:] + subNbr + ".jpg"
        # taxe
        elif type_prefix == "TA":
            nbr = "T01-270-" + stampNumber_str[3:] + subNbr + ".jpg"
        # croix rouge
        elif type_prefix == "CR":
            nbr = "T01-565-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Guerre
        elif type_prefix == "GU":
            nbr = "T01-150-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Alsace Lorraine
        elif type_prefix == "AL":
            nbr = "T01-110-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Journaux
        elif type_prefix == "JN":
            nbr = "T01-160-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Petit colis
        elif type_prefix == "PC":
            nbr = "T01-120-" + stampNumber_str[3:] + subNbr + ".jpg"
        # colis postaux
        elif type_prefix == "CP":
            nbr = "T01-130-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Fiscaux postaux
        elif type_prefix == "FP":
            nbr = "T01-140-" + stampNumber_str[3:] + subNbr + ".jpg"
        # CFA
        elif type_prefix == "CF":
            nbr = "T01-180-CF-" + stampNumber_str[3:] + subNbr + ".jpg"
        # colonies
        elif type_prefix == "CO":
            nbr = "T01-200-" + stampNumber_str[3:] + subNbr + ".jpg"
        # télégraphe
        elif type_prefix == "TL":
            nbr = "T01-210-" + stampNumber_str[3:] + subNbr + ".jpg"
        # telegraphe militaire
        elif type_prefix == "MT":
            nbr = "T01-211-" + stampNumber_str[3:] + subNbr + ".jpg"
        # téléphone
        elif type_prefix == "TE":
            nbr = "T01-215-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Greve
        elif type_prefix == "GR":
            nbr = "T01-220-" + stampNumber_str[3:] + subNbr + ".jpg"
        # Poste Montenegrine
        elif type_prefix == "PM":
            nbr = "T01-230-" + stampNumber_str[3:] + subNbr + ".jpg"
        # fictifs
        elif type_prefix == "F ":
            nbr = "T01-240-" + stampNumber_str[2:] + subNbr + ".jpg"
        # fictifs
        elif type_prefix == "F-":
            nbr = "T01-000-F" + stampNumber_str[2:] + subNbr + ".jpg"
        # bienfaisance
        elif type_prefix == "BI":
            nbr = "T01-250-" + stampNumber_str[2:] + subNbr + ".jpg"
        # mandat
        elif type_prefix == "MA":
            nbr = "T01-251-" + stampNumber_str[2:] + subNbr + ".jpg"
        # distributeur
        elif type_prefix == "DI":
            nbr = "T01-260-" + stampNumber_str[2:] + subNbr + ".jpg"
        else:
            stampType_str = str(stampType) if stampType else ""
            if stampType_str == "COURS INSTRUCTION":
                nbr = "T01-252-" + stampNumber_str + subNbr + ".jpg"
            elif stampType_str == "Préoblitéré COURS INSTRUCTION":
                nbr = "T01-253-" + stampNumber_str + subNbr + ".jpg"
            elif stampType_str == "RECOUVREMENT COURS INSTRUCTION":
                nbr = "T01-254-" + stampNumber_str + subNbr + ".jpg"
            elif stampType_str.strip() == "TAXES COURS INSTRUCTION":
                nbr = "T01-255-" + stampNumber_str + subNbr + ".jpg"
            elif stampType_str.strip() == "POSTE AERIENNE MILITAIRE":
                nbr = "T01-256-" + stampNumber_str + subNbr + ".jpg"
            else:
                nbr = "T01-000-" + stampNumber_str + subNbr + ".jpg"

        return nbr

    def _load_stamp_image(self, stamp_nbr, stamp_type, stamp_sub_nbr=""):
        """Display stamp photo using exact getPhotoName and scaling rules."""
        rel_filename = self.getPhotoName(stamp_nbr, stamp_type, stamp_sub_nbr)

        # Build potential absolute image paths
        possible_paths = [
            self.base_dir / "images" / self.current_country / rel_filename,
            self.base_dir / "photos" / self.current_country / rel_filename,
            self.base_dir / "images" / rel_filename,
            self.base_dir / "photos" / rel_filename,
            self.base_dir / rel_filename
        ]

        found_path = next((p for p in possible_paths if p.exists()), None)

        if found_path:
            pixmap = QPixmap(str(found_path))
            if not pixmap.isNull():
                if pixmap.width() > pixmap.height():
                    scaled_pixmap = pixmap.scaledToWidth(200)
                else:
                    scaled_pixmap = pixmap.scaledToHeight(200)

                self.img_label.setPixmap(scaled_pixmap)
                return

        self.img_label.setText(_("No Photo Available"))

    def _save_stamp_status(self):
        if not self.selected_nbr or self.current_collection_id is None:
            return

        status = self.combo_status.currentText()
        if status == _("Not Owned"):
            self.central_db.remove_item(self.current_collection_id, self.selected_nbr)
        else:
            self.central_db.save_item(
                collection_id=self.current_collection_id,
                stamp_nbr=self.selected_nbr,
                status=status,
                condition=self.combo_condition.currentText(),
                quantity=self.spin_qty.value(),
                price=self.spin_price.value(),
                notes=self.txt_notes.toPlainText().strip()
            )

        self._refresh_table()

    def _export_mancolist(self):
        """Export missing stamps for the active collection to Excel."""
        if not self.country_db or self.current_collection_id is None:
            return

        stamps = self.country_db.get_stamps()
        user_items = self.central_db.get_collection_items(self.current_collection_id)

        mancolist = []
        for stamp in stamps:
            nbr, year, stype, valcolor, desc, key, sub_nbr = stamp
            user_data = user_items.get(str(nbr))
            status = user_data[0] if user_data else _("Missing")

            if not user_data or status in [_("Manco (Wanted)"), _("Missing")]:
                notes = user_data[4] if user_data else ""
                mancolist.append((nbr, year, stype, valcolor, desc, status, notes))

        if not mancolist:
            QMessageBox.information(self, _("Information"), _("No missing stamps found for this collection."))
            return

        coll_name = self.combo_collections.currentText().split(" (")[0]
        default_filename = f"Mancolist_{coll_name}.xlsx"

        file_path, _filter = QFileDialog.getSaveFileName(
            self, _("Save Mancolist"), default_filename, "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _("Mancolist")

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1D6F42", end_color="1D6F42", fill_type="solid")

        headers = [
            _("Yvert N°"),
            _("Year"),
            _("Type"),
            _("Value / Color"),
            _("Description"),
            _("Status"),
            _("Notes")
        ]
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_data in mancolist:
            ws.append(list(row_data))

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(file_path)
        QMessageBox.information(self, _("Success"), _("Mancolist exported successfully:\n") + file_path)

    def _clear_ui(self):
        self.table.setRowCount(0)
        self.lbl_valeur.setText("-")
        self.lbl_desc.setText("-")
        self.img_label.setText(_("No Photo Available"))

    def closeEvent(self, event):
        if self.central_db:
            self.central_db.close()
        if self.country_db:
            self.country_db.close()
        event.accept()


if __name__ == '__main__':
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())